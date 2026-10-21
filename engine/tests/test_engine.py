"""Engine tests on generated fixtures (made-up people only).

The Worker normally decides what to replace; here a tiny stand-in finds the
real strings in the extracted text so each format's rewrite can be checked.
"""

from __future__ import annotations

import io
import json
import os
import re
import zipfile

import pymupdf as fitz
import pytest
from docx import Document
from fastapi.testclient import TestClient
from lxml import etree

from engine.main import DOCX_MIME, PDF_MIME, app

client = TestClient(app)

LINES = [
    "Assurances Mirador inc.",
    "Avis de renouvellement - Assurance habitation",
    "",
    "Assuré : M. Jean-François Tremblay",
    "Coassurée : Mme Isabelle Côté",
    "Adresse du risque : 1842, rue des Érables, Laval (Québec) H7N 2K4",
    "Téléphone : (450) 662-1934   Courriel : jf.tremblay@videotron.example",
    "Numéro de police : HAB-48210937",
    "",
    "Monsieur Tremblay, votre police sera renouvelée le 1er décembre 2026.",
    "Signé : J.-F. Tremblay et Isabelle Côté",
]

SWAPS = {
    "Jean-François Tremblay": "Pierre-Luc Larivière",
    "J.-F. Tremblay": "P.-L. Larivière",
    "Tremblay": "Larivière",
    "Isabelle Côté": "Solange Guay",
    "Côté": "Guay",
    "1842, rue des Érables": "1238, rue des Tilleuls",
    "1842 rue des Érables": "1238 rue des Tilleuls",
    "H7N 2K4": "H7J 2H4",
    "662-1934": "555-0142",
    "jf.tremblay@videotron.example": "pl.lariviere@example.com",
    "HAB-48210937": "HAB-77882353",
}
REAL = ["Tremblay", "Côté", "Érables", "662-1934", "H7N 2K4", "48210937", "videotron"]


def plan(text: str, swaps: dict[str, str] = SWAPS) -> list[dict]:
    """Longest-first, non-overlapping literal matches: a stand-in for the Worker's matcher."""
    taken: list[tuple[int, int]] = []
    out = []
    for real in sorted(swaps, key=len, reverse=True):
        for m in re.finditer(re.escape(real), text):
            a, b = m.span()
            if any(a < y and b > x for x, y in taken):
                continue
            taken.append((a, b))
            out.append({"start": a, "end": b, "text": swaps[real]})
    return sorted(out, key=lambda r: r["start"])


def apply(text: str, reps: list[dict]) -> str:
    out, at = "", 0
    for r in reps:
        out += text[at : r["start"]] + r["text"]
        at = r["end"]
    return out + text[at:]


def squash(s: str) -> str:
    return re.sub(r"\s+", "", s)


def make_pdf() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    y = 72
    for i, line in enumerate(LINES):
        font = "tiro" if i % 2 else "helv"
        page.insert_text((72, y), line, fontsize=11, fontname=font)
        y += 18
    page.add_text_annot((400, 60), "Appeler Jean-François Tremblay")
    page.insert_link({"kind": fitz.LINK_URI, "from": fitz.Rect(72, 170, 300, 182), "uri": "mailto:jf.tremblay@videotron.example"})
    doc.set_toc([[1, "Dossier Tremblay", 1]])
    doc.set_metadata({"author": "Jean-François Tremblay", "title": "Assurance Tremblay"})
    return doc.tobytes()


def make_scan(pdf_bytes: bytes, turn: int = 0) -> bytes:
    """A 200 dpi picture of the first page; `turn` rotates the pixels clockwise, like a sheet fed sideways."""
    src = fitz.open(stream=pdf_bytes, filetype="pdf")
    pix = src[0].get_pixmap(matrix=fitz.Matrix(200 / 72, 200 / 72).prerotate(turn), alpha=False)
    out = fitz.open()
    page = out.new_page(width=pix.width * 72 / 200, height=pix.height * 72 / 200)
    page.insert_image(page.rect, stream=pix.tobytes("png"))
    return out.tobytes()


def post_rewrite(data: bytes, mime: str, sha: str, reps: list[dict], fields: list[dict] | None = None):
    return client.post(
        "/rewrite",
        files={"file": ("f", data, mime)},
        data={"spec": json.dumps({"sha256": sha, "replacements": reps, "fields": fields or []}), "mime": mime},
    )


def test_text_pdf_rewrite_in_place():
    data = make_pdf()
    ex = client.post("/extract", content=data, headers={"x-mime": PDF_MIME}).json()
    assert "Jean-François Tremblay" in ex["text"]
    reps = plan(ex["text"])
    field_ids = {f["id"]: f["text"] for f in ex["fields"]}
    assert any("Tremblay" in t for t in field_ids.values())
    new_fields = [{"id": k, "text": apply(v, plan(v))} for k, v in field_ids.items()]

    r = post_rewrite(data, PDF_MIME, ex["sha256"], reps, new_fields)
    assert r.status_code == 200, r.text
    check = client.post("/check", content=r.content, headers={"x-mime": PDF_MIME}).json()
    for s in REAL:
        assert s not in check["text"], s
        assert all(s not in f["text"] for f in check["fields"]), s
        assert all(s not in page for page in check["raster"]), s
    assert not check.get("metadata")
    # Nothing else on the page moved or disappeared.
    assert squash(check["text"]) == squash(apply(ex["text"], reps))
    # A plain text page isn't OCR'd again: its picture shows nothing its text layer doesn't (pdf.regions_to_read).
    assert check["raster"] == [""]
    # Replacements sit where the originals were.
    out = fitz.open(stream=r.content, filetype="pdf")
    hits = out[0].search_for("Larivière")
    assert hits and all(72 <= h.x0 < 400 for h in hits)


def test_rejects_stale_ranges():
    data = make_pdf()
    r = post_rewrite(data, PDF_MIME, "0" * 64, [{"start": 0, "end": 3, "text": "x"}])
    assert r.status_code == 409


def visible_spans(page: fitz.Page, needle: str) -> list[dict]:
    """Drawn (not invisible OCR layer) text spans holding `needle`."""
    return [s for s in page.get_texttrace() if s["type"] == 0 and needle in "".join(chr(c[0]) for c in s["chars"])]


def test_scanned_pdf_ocr_and_patch():
    scan = make_scan(make_pdf())
    info = client.post("/inspect", content=scan, headers={"x-mime": PDF_MIME}).json()
    assert info["ocr"] == "scan"
    ocred = client.post("/ocr", content=scan, headers={"x-mime": PDF_MIME})
    assert ocred.status_code == 200, ocred.text
    ex = client.post("/extract", content=ocred.content, headers={"x-mime": PDF_MIME}).json()
    # The text layer is invisible text over the untouched picture, so pixels get patched.
    assert ex["pageKinds"] == ["ocr"]
    # Every word read right; the note's icon is in the picture too, and reads "[=]".
    assert squash(ex["text"]).replace("[=]", "") == squash("\n".join(LINES))
    src = fitz.open(stream=scan, filetype="pdf")
    out = fitz.open(stream=ocred.content, filetype="pdf")
    assert fitz.Pixmap(out, out[0].get_images()[0][0]).samples == fitz.Pixmap(src, src[0].get_images()[0][0]).samples
    reps = plan(ex["text"])
    r = post_rewrite(ocred.content, PDF_MIME, ex["sha256"], reps)
    assert r.status_code == 200, r.text
    check = client.post("/check", content=r.content, headers={"x-mime": PDF_MIME}).json()
    raster = "\n".join(check["raster"])
    for s in ["Tremblay", "662-1934", "48210937"]:
        assert s not in check["text"], s
        assert s not in raster, f"{s} still visible in the pixels"
    assert "Larivière" in raster
    # Replacements are drawn on the original baseline, at about the original size (11 pt).
    page = fitz.open(stream=r.content, filetype="pdf")[0]
    [span] = visible_spans(page, "Pierre-Luc")
    assert abs(span["chars"][0][2][1] - (72 + 3 * 18)) < 0.75, span["chars"][0][2]
    assert 10 <= span["size"] <= 12, span["size"]


def test_rotated_scan_is_turned_upright():
    for turn in (90, 180):
        scan = make_scan(make_pdf(), turn)
        ocred = client.post("/ocr", content=scan, headers={"x-mime": PDF_MIME})
        assert ocred.status_code == 200, ocred.text
        page = fitz.open(stream=ocred.content, filetype="pdf")[0]
        assert page.rotation == 0 and page.rect.width < page.rect.height, (turn, page.rect)
        ex = client.post("/extract", content=ocred.content, headers={"x-mime": PDF_MIME}).json()
        assert ex["pageKinds"] == ["ocr"]
        assert "Jean-François Tremblay" in ex["text"] and "HAB-48210937" in ex["text"], (turn, ex["text"])
        r = post_rewrite(ocred.content, PDF_MIME, ex["sha256"], plan(ex["text"]))
        assert r.status_code == 200, r.text
        check = client.post("/check", content=r.content, headers={"x-mime": PDF_MIME}).json()
        raster = "\n".join(check["raster"])
        assert "Tremblay" not in raster and "Larivière" in raster, (turn, raster)


def test_only_scanned_pages_get_a_text_layer():
    doc = fitz.open(stream=make_pdf(), filetype="pdf")
    doc.insert_pdf(fitz.open(stream=make_scan(make_pdf()), filetype="pdf"))
    data = doc.tobytes()
    assert client.post("/inspect", content=data, headers={"x-mime": PDF_MIME}).json()["pageKinds"] == ["text", "scan"]
    r = client.post("/ocr", content=data, headers={"x-mime": PDF_MIME})
    assert r.status_code == 200, r.text
    out = fitz.open(stream=r.content, filetype="pdf")
    assert client.post("/inspect", content=r.content, headers={"x-mime": PDF_MIME}).json()["pageKinds"] == ["text", "ocr"]
    assert out[0].read_contents() == doc[0].read_contents()
    assert out[0].get_text() == doc[0].get_text()
    assert "Jean-François Tremblay" in out[1].get_text()


def outlined_page(doc: fitz.Document) -> None:
    """A page whose text is drawn as shapes, like some "print to PDF" exports: no text to extract, hundreds of paths."""
    src = fitz.open(stream=make_pdf(), filetype="pdf")[0]
    pix = src.get_pixmap(dpi=100, colorspace=fitz.csGRAY, clip=fitz.Rect(0, 0, src.rect.width, 260))
    page = doc.new_page(width=src.rect.width, height=src.rect.height)
    k = 72 / 100
    shape = page.new_shape()
    for y in range(pix.height):
        x = 0
        while x < pix.width:
            if pix.pixel(x, y)[0] < 128:
                x0 = x
                while x < pix.width and pix.pixel(x, y)[0] < 128:
                    x += 1
                shape.draw_rect(fitz.Rect(x0 * k, y * k, x * k, (y + 1) * k))
            x += 1
    shape.finish(color=None, fill=(0, 0, 0))
    shape.commit()


def test_force_ocr_rasterizes_only_unreadable_pages_and_keeps_annotations():
    doc = fitz.open(stream=make_pdf(), filetype="pdf")
    outlined_page(doc)
    data = doc.tobytes()
    assert client.post("/inspect", content=data, headers={"x-mime": PDF_MIME}).json()["pageKinds"] == ["text", "garbled"]
    r = client.post("/ocr", content=data, headers={"x-mime": PDF_MIME, "x-force": "1"})
    assert r.status_code == 200, r.text
    out = fitz.open(stream=r.content, filetype="pdf")
    # The readable page keeps its fonts, links and annotation.
    assert "OCRText" not in [f[4] for f in out[0].get_fonts()]
    assert out[0].first_annot is not None and out[0].get_links()
    # The outlined one is now a picture with an OCR layer.
    assert [f[4] for f in out[1].get_fonts()] == ["OCRText"] and len(out[1].get_images()) == 1
    ex = client.post("/extract", content=r.content, headers={"x-mime": PDF_MIME}).json()
    assert ex["pageKinds"] == ["text", "ocr"]
    # The outline page is blocky (drawn from 100 dpi pixels); OCR still reads the name on it.
    page2 = ex["text"].split("Assurances Mirador")[2]
    assert "Tremblay" in page2 and "Côté" in page2, page2


def test_image_becomes_pdf():
    pdf_bytes = make_pdf()
    pix = fitz.open(stream=pdf_bytes, filetype="pdf")[0].get_pixmap(dpi=150)
    r = client.post("/ocr", content=pix.tobytes("jpeg"), headers={"x-mime": "image/jpeg"})
    assert r.status_code == 200
    ex = client.post("/extract", content=r.content, headers={"x-mime": PDF_MIME}).json()
    assert "Tremblay" in ex["text"]


def test_photo_orientation_and_heic():
    from PIL import Image

    pix = fitz.open(stream=make_pdf(), filetype="pdf")[0].get_pixmap(dpi=150)
    upright = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    # A phone stores the sensor's pixels and a tag saying to turn them 90° clockwise (EXIF orientation 6).
    exif = Image.Exif()
    exif[0x0112] = 6
    buf = io.BytesIO()
    upright.transpose(Image.Transpose.ROTATE_90).save(buf, format="JPEG", quality=90, exif=exif.tobytes())
    jpeg = buf.getvalue()
    r = client.post("/ocr", content=jpeg, headers={"x-mime": "image/jpeg"})
    assert r.status_code == 200, r.text
    out = fitz.open(stream=r.content, filetype="pdf")
    page = out[0]
    assert page.rect.width < page.rect.height
    assert out.extract_image(page.get_images()[0][0])["image"] == jpeg  # the camera's bytes, as they were
    assert "Jean-François Tremblay" in client.post("/extract", content=r.content, headers={"x-mime": PDF_MIME}).json()["text"]

    buf = io.BytesIO()
    upright.save(buf, format="HEIF", quality=90)
    r = client.post("/ocr", content=buf.getvalue(), headers={"x-mime": "image/heic"})
    assert r.status_code == 200, r.text
    assert "Jean-François Tremblay" in client.post("/extract", content=r.content, headers={"x-mime": PDF_MIME}).json()["text"]


def test_nothing_is_written_to_disk(tmp_path, monkeypatch):
    """OCR, rewrite and check work in memory: no temp file, not even an upload spooled past 1 MB."""
    import tempfile

    made: list[str] = []
    for name in ("TemporaryFile", "NamedTemporaryFile", "mkstemp", "mkdtemp"):
        real = getattr(tempfile, name)
        monkeypatch.setattr(tempfile, name, lambda *a, _n=name, _r=real, **k: (made.append(_n), _r(*a, **k))[1])
    monkeypatch.setattr(tempfile, "tempdir", str(tmp_path))
    monkeypatch.setenv("TMPDIR", str(tmp_path))

    ocred = client.post("/ocr", content=make_scan(make_pdf()), headers={"x-mime": PDF_MIME}).content
    ex = client.post("/extract", content=ocred, headers={"x-mime": PDF_MIME}).json()
    r = post_rewrite(ocred, PDF_MIME, ex["sha256"], plan(ex["text"]))
    assert r.status_code == 200, r.text
    assert client.post("/check", content=r.content, headers={"x-mime": PDF_MIME}).status_code == 200

    pix = fitz.open(stream=make_pdf(), filetype="pdf")[0].get_pixmap(dpi=150)
    assert client.post("/ocr", content=pix.tobytes("png"), headers={"x-mime": "image/png"}).status_code == 200

    big = fitz.open(stream=make_pdf(), filetype="pdf")
    big.embfile_add("ballast.bin", os.urandom(3 * 1024 * 1024))
    big = big.tobytes()
    ex = client.post("/extract", content=big, headers={"x-mime": PDF_MIME}).json()
    r = post_rewrite(big, PDF_MIME, ex["sha256"], plan(ex["text"]))
    assert r.status_code == 200, r.text

    assert made == []
    assert list(tmp_path.iterdir()) == []


def test_health_answers_while_a_document_is_processed(monkeypatch):
    import threading
    import time

    from engine import main

    started, release = threading.Event(), threading.Event()
    busy, peak = [0], [0]

    def slow_ocr(data, force):
        busy[0] += 1
        peak[0] = max(peak[0], busy[0])
        started.set()
        release.wait(10)
        busy[0] -= 1
        return data

    monkeypatch.setattr(main.ocr, "ocr_pdf", slow_ocr)
    data = make_pdf()
    # One client, so one event loop: a handler that blocked it would hold /health up too.
    with TestClient(app) as tc:
        codes: list[int] = []
        jobs = [threading.Thread(target=lambda: codes.append(tc.post("/ocr", content=data, headers={"x-mime": PDF_MIME}).status_code)) for _ in range(2)]
        for j in jobs:
            j.start()
        assert started.wait(10)
        t0 = time.perf_counter()
        assert tc.get("/health").status_code == 200
        assert time.perf_counter() - t0 < 1.0
        assert not release.is_set() and busy[0] == 1
        release.set()
        for j in jobs:
            j.join(10)
    assert codes == [200, 200]
    assert peak[0] == 1  # documents are handled one at a time


def make_docx() -> bytes:
    d = Document()
    d.core_properties.author = "Jean-François Tremblay"
    d.core_properties.title = "Dossier Tremblay"
    d.sections[0].header.paragraphs[0].text = "Confidentiel - Isabelle Côté"
    p = d.add_paragraph("Assuré : ")
    for piece in ["Jean-", "Fran", "çois ", "Trem", "blay"]:
        r = p.add_run(piece)
        r.bold = piece.startswith("Trem")
    d.add_paragraph("Adresse : 1842, rue des Érables, Laval")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Police"
    t.cell(0, 1).text = "HAB-48210937"
    d.add_paragraph("Contact : jf.tremblay@videotron.example")
    buf = io.BytesIO()
    d.save(buf)
    # Add a tracked deletion that still holds the old phone number.
    z = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    files = {n: z.read(n) for n in z.namelist()}
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    tree = etree.fromstring(files["word/document.xml"])
    body = tree.find(f"{{{W}}}body")
    p = etree.SubElement(body, f"{{{W}}}p")
    dl = etree.SubElement(p, f"{{{W}}}del", {f"{{{W}}}author": "Jean-François Tremblay", f"{{{W}}}id": "1", f"{{{W}}}date": "2026-01-01T00:00:00Z"})
    r = etree.SubElement(dl, f"{{{W}}}r")
    etree.SubElement(r, f"{{{W}}}delText").text = "Ancien tél. 450-662-1934"
    body.remove(p)
    body.insert(0, p)
    files["word/document.xml"] = etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zz:
        for n, b in files.items():
            zz.writestr(n, b)
    return out.getvalue()


def test_docx_rewrite_across_runs_headers_tables_and_history():
    data = make_docx()
    ex = client.post("/extract", content=data, headers={"x-mime": DOCX_MIME}).json()
    assert "Jean-François Tremblay" in ex["text"]
    assert "662-1934" not in ex["text"]  # deleted text isn't content…
    reps = plan(ex["text"])
    r = post_rewrite(data, DOCX_MIME, ex["sha256"], reps)
    assert r.status_code == 200, r.text
    raw = b"".join(zipfile.ZipFile(io.BytesIO(r.content)).read(n) for n in zipfile.ZipFile(io.BytesIO(r.content)).namelist())
    text = raw.decode("utf-8", "replace")
    for s in REAL:
        assert s not in text, f"{s} still somewhere in the package"  # …and it's gone from the file too
    check = client.post("/check", content=r.content, headers={"x-mime": DOCX_MIME}).json()
    assert squash(check["text"]) == squash(apply(ex["text"], reps))
    out = Document(io.BytesIO(r.content))
    assert out.core_properties.author == ""
    assert out.sections[0].header.paragraphs[0].text == "Confidentiel - Solange Guay"
    para = next(p for p in out.paragraphs if p.text.startswith("Assuré"))
    assert para.text == "Assuré : Pierre-Luc Larivière"


def test_subset_fonts_are_reused_and_reading_order_kept():
    """Real-world PDFs embed subset fonts with no Unicode map; the page must still be retypeset in its own fonts."""
    from pathlib import Path

    from engine import pdf as pdfmod

    path = Path(__file__).resolve().parents[2] / "fixtures" / "files" / "assurance-habitation.pdf"
    if not path.exists():
        pytest.skip("run fixtures/make_files.py first")
    data = path.read_bytes()
    doc = pdfmod.open_pdf(data)
    ex = pdfmod.extract(doc)
    reps = plan(ex.text)
    stats = pdfmod.rewrite(doc, ex.sha256, reps, [], set())
    assert stats["retypeset_pages"] == 1, stats
    out = fitz.open(stream=pdfmod.save(doc), filetype="pdf")
    text = out[0].get_text()
    for s in REAL:
        assert s not in text, s
    assert squash(text) == squash(apply(ex.text, reps))
    # The page's own (subset) fonts drew the unchanged text.
    fonts = {f[3].split("+")[-1] for f in out[0].get_fonts()}
    assert any("Liberation Sans" in f for f in fonts), fonts


def test_pages_are_read_in_parallel_with_the_same_result(monkeypatch):
    """Three scanned pages, one upside down: workers or not, the same text layer and the same check."""
    pages = [fitz.open(stream=make_scan(make_pdf(), turn), filetype="pdf") for turn in (0, 180, 0)]
    doc = fitz.open()
    for p in pages:
        doc.insert_pdf(p)
    scan = doc.tobytes()

    def run(workers: str):
        monkeypatch.setenv("ANON_OCR_WORKERS", workers)
        ocred = client.post("/ocr", content=scan, headers={"x-mime": PDF_MIME})
        assert ocred.status_code == 200, ocred.text
        ex = client.post("/extract", content=ocred.content, headers={"x-mime": PDF_MIME}).json()
        out = post_rewrite(ocred.content, PDF_MIME, ex["sha256"], plan(ex["text"]))
        check = client.post("/check", content=out.content, headers={"x-mime": PDF_MIME}).json()
        return ex["text"], check["raster"]

    one = run("1")
    three = run("3")
    assert one == three
    assert one[0].count("Jean-François Tremblay") == 3, one[0]
    # The check reads around each patch: the fake is there, the real name isn't.
    assert len(one[1]) == 3 and all("Larivi" in r and "Tremblay" not in r for r in one[1]), one[1]


def test_check_reads_a_name_in_an_image_on_a_text_page():
    """Only the image's area is OCR'd, and the name in it is found."""
    doc = fitz.open(stream=make_pdf(), filetype="pdf")
    stamp = fitz.open()
    sp = stamp.new_page(width=300, height=40)
    sp.insert_text((10, 28), "Signé : Isabelle Côté", fontsize=18)
    png = sp.get_pixmap(dpi=200).tobytes("png")
    doc[0].insert_image(fitz.Rect(300, 600, 500, 627), stream=png)
    data = doc.tobytes()
    check = client.post("/check", content=data, headers={"x-mime": PDF_MIME}).json()
    assert "Côté" in check["raster"][0] or "Cote" in check["raster"][0], check["raster"]
    # The page's own text isn't read again.
    assert "Mirador" not in check["raster"][0]
